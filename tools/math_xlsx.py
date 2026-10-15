#!/usr/bin/env python3
"""Таблица математики для заказчика: math/Senbazuru_MATH_v3.xlsx.

    python3 tools/math_xlsx.py <verify.json от tools/verify_books.py> [выход.xlsx]

Листы: режимы (RTP, max win, hit rate, бонус, потолок, σ), выплаты, параметры движка, страты выгрузки.
"""
import sys, json, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ver = json.load(open(sys.argv[1]))
out = sys.argv[2] if len(sys.argv) > 2 else 'math/Senbazuru_MATH_v3.xlsx'
rep = json.load(open('math/stake_report.json'))
cfg = json.loads(subprocess.check_output(['node', '-e', "const M=require('./assets/js/math.js');console.log(JSON.stringify({c:M.CONFIG,t:M.TUNE}))"]))
NAMES = {'base': 'Обычное вращение', 'ante2': 'Анте ×2', 'ante4': 'Анте ×4', 'bonus': 'Покупка: фриспины', 'super': 'Покупка: суперфриспины',
         'lantern': 'Покупка: фонари линий', 'fold': 'Покупка: сгиб'}
HEAD, BOLD = PatternFill('solid', fgColor='17224D'), Font(bold=True, color='FFFFFF')

def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]: c.fill, c.font, c.alignment = HEAD, BOLD, Alignment(wrap_text=True, vertical='center')
    for r in rows: ws.append(r)
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    return ws

wb = Workbook(); wb.remove(wb.active)
rows = []
for k, v in ver.items():
    m = rep['modes'][k]
    rows.append([k, NAMES.get(k, k), v['cost'], v['books'], round(v['rtp'] * 100, 4), round(m['model']['rtp'] * 100, 3), v['maxWin'], round(v['hitRate'] * 100, 2),
                 round(1 / v['bonusRate'], 1) if v['bonusRate'] < 0.999 else 1, round(1 / v['capRate']) if v['capRate'] else None, round(v['sdPerCost'], 2), v['problems']])
ws = sheet(wb, 'Режимы', ['Режим RGS', 'Название', 'Цена, × ставки', 'Книг', 'RTP таблицы, %', 'RTP движка (оценка), %', 'Max win, × ставки',
                          'Hit rate, %', 'Бонус: 1 из N раундов', 'Потолок 15 000×: 1 из N', 'σ выигрыша / цена', 'Ошибок пересчёта'], rows,
           [11, 26, 10, 10, 12, 14, 12, 10, 14, 16, 12, 12])
ws.append([]); ws.append(['RTP таблицы = Σ(вес · выплата / 100) / Σ вес / цена по lookUpTable_<режим>_0.csv. Пересчёт книг по доскам — tools/verify_books.py.'])

pt = cfg['c']['paytable']
sheet(wb, 'Выплаты', ['Символ', '3 на линии', '4 на линии', '5 на линии'], [[s] + [x / 100 for x in pt[s]] for s in pt], [10, 12, 12, 12]).append(
    ['× ставки; 5 линий по рядам, слева направо, платит старшая комбинация линии'])

t = cfg['t']
prm = [
    ['Поле', '5×5, 5 линий (по ряду)'], ['Потолок выигрыша', f"{cfg['c']['maxWin']}× ставки, раунд заканчивается"],
    ['Вайлд', 'барабаны 2–4'], ['Множители вайлда, база', str(t['base']['wildMult'])], ['Множители вайлда, фриспины', str(t['fs']['wildMult'])],
    ['Фриспины за 3/4/5 скаттеров', str(cfg['c']['fsAward'])], ['Ретриггер', '3+ скаттера: +5, всего не больше 40'],
    ['Скаттер на барабан: база', t['base']['scatter']], ['Скаттер на барабан: анте ×2', t['ante2']['scatter']], ['Скаттер на барабан: анте ×4', t['ante4']['scatter']],
    ['Фонари линий: прибавки вайлда (вес)', str(t['lantern']['wildMult'])], ['Фонари линий: вес вайлда / скаттер', f"{t['lantern']['wild']} / {t['lantern']['scatter']}"],
    ['Фонари линий: ретриггер', '+5 вращений и ×2 ко всем фонарям'],
    ['Сгиб: множители вайлда (вес)', str(t['fold']['wildMult'])], ['Сгиб: вес вайлда', t['fold']['wild']], ['Сгиб: отпечаток вайлда', '×2 к множителю; скаттеров нет'],
    ['Суперфриспины: вес вайлда', t['super']['wild']], ['Веса символов L1…H1', str(t['base']['sym'])],
]
sheet(wb, 'Параметры', ['Параметр', 'Значение'], prm, [38, 60])

srows = []
for k, m in rep['modes'].items():
    for s in m['strata']: srows.append([k, s['label'], s['prob'], s['books'], round(s['rtpShare'] * 100, 4), round(s['meanBooks'], 3)])
sheet(wb, 'Страты', ['Режим', 'Страта', 'Вероятность', 'Книг', 'Вклад в RTP, %', 'Средний выигрыш в книгах, ×'], srows, [10, 40, 14, 10, 14, 16])
wb.save(out); print('xlsx:', out)
